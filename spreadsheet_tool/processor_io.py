from __future__ import annotations

import os
import shutil
from pathlib import Path

import pandas as pd

from .models import ExportSettings, SourceSelection
from .processor_common import is_empty_value
from .processor_pipeline import cell_display_length, prepare_dataframe_for_excel

AUTO_WIDTH_SAMPLE_LIMIT = 2000


def export_dataframe(
    dataframe: pd.DataFrame,
    output_path: str | Path,
    settings: ExportSettings,
) -> None:
    path = Path(output_path)
    if settings.output_format == "csv":
        dataframe.to_csv(path, index=False, encoding="utf-8-sig")
        return
    export_to_excel(dataframe, path, settings)


def export_dataframe_with_old_workbook(
    dataframe: pd.DataFrame,
    source: SourceSelection,
    output_path: str | Path,
    settings: ExportSettings,
) -> None:
    source_path = source.path
    if source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持基于 Excel 老文件导出整本工作簿。")

    target_path = Path(output_path)
    if paths_refer_to_same_file(source_path, target_path):
        raise ValueError("完整老文件导出不能覆盖原文件，请选择新的保存位置。")

    shutil.copy2(source_path, target_path)
    write_dataframe_to_existing_excel_sheet(dataframe, target_path, source.sheet_name, settings)


def export_to_excel(
    dataframe: pd.DataFrame,
    output_path: Path,
    settings: ExportSettings,
) -> None:
    excel_ready = prepare_dataframe_for_excel(dataframe)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        excel_ready.to_excel(writer, index=False, sheet_name=settings.sheet_name)
        worksheet = writer.book[settings.sheet_name]
        if settings.freeze_header:
            worksheet.freeze_panes = "A2"
        if settings.style_header:
            apply_header_style(worksheet)
        if settings.auto_width:
            apply_auto_width(worksheet, dataframe)


def write_dataframe_back_to_source(
    dataframe: pd.DataFrame,
    source: SourceSelection,
    settings: ExportSettings,
) -> None:
    path = source.path
    suffix = path.suffix.lower()

    if suffix == ".csv":
        dataframe.to_csv(path, index=False, encoding="utf-8-sig")
        return
    if suffix == ".tsv":
        dataframe.to_csv(path, index=False, sep="\t", encoding="utf-8-sig")
        return
    if suffix in {".xlsx", ".xlsm"}:
        write_dataframe_to_existing_excel_sheet(dataframe, path, source.sheet_name, settings)
        return

    raise ValueError(f"暂不支持回写到该文件类型: {path.name}")


def write_dataframe_to_existing_excel_sheet(
    dataframe: pd.DataFrame,
    output_path: Path,
    sheet_name: str,
    settings: ExportSettings,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils.dataframe import dataframe_to_rows

    workbook = load_workbook(output_path, keep_vba=output_path.suffix.lower() == ".xlsm")
    excel_ready = prepare_dataframe_for_excel(dataframe)
    try:
        if sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            clear_worksheet_values_preserve_structure(worksheet, MergedCell)
            unmerge_ranges_overlapping_write_area(worksheet, len(excel_ready) + 1, len(excel_ready.columns))
            write_rows_to_worksheet(worksheet, excel_ready, dataframe_to_rows)
        else:
            worksheet = workbook.create_sheet(title=sheet_name, index=len(workbook.sheetnames))
            write_rows_to_worksheet(worksheet, excel_ready, dataframe_to_rows)
            if settings.freeze_header:
                worksheet.freeze_panes = "A2"
            if settings.style_header and dataframe.shape[1] > 0:
                apply_header_style(worksheet)
            if settings.auto_width and dataframe.shape[1] > 0:
                apply_auto_width(worksheet, dataframe)
        workbook.save(output_path)
    finally:
        workbook.close()


def clear_worksheet_values_preserve_structure(worksheet: object, merged_cell_type: type[object]) -> None:
    max_row = getattr(worksheet, "max_row", 0) or 0
    max_column = getattr(worksheet, "max_column", 0) or 0
    if max_row <= 0 or max_column <= 0:
        return
    for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
        for cell in row:
            if isinstance(cell, merged_cell_type):
                continue
            cell.value = None


def unmerge_ranges_overlapping_write_area(worksheet: object, max_row: int, max_column: int) -> None:
    if max_row <= 0 or max_column <= 0:
        return
    for merged_range in list(worksheet.merged_cells.ranges):
        if (
            merged_range.min_row <= max_row
            and merged_range.max_row >= 1
            and merged_range.min_col <= max_column
            and merged_range.max_col >= 1
        ):
            worksheet.unmerge_cells(str(merged_range))


def write_rows_to_worksheet(worksheet: object, dataframe: pd.DataFrame, dataframe_to_rows_func: object) -> None:
    for row_index, row in enumerate(dataframe_to_rows_func(dataframe, index=False, header=True), start=1):
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)


def apply_header_style(worksheet: object) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    font = Font(bold=True)
    alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment


def apply_auto_width(worksheet: object, dataframe: pd.DataFrame | None = None) -> None:
    from openpyxl.utils import get_column_letter

    if dataframe is None:
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            max_length = 0
            for cell in column_cells[:AUTO_WIDTH_SAMPLE_LIMIT]:
                text = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(text))
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 40)
        return

    sample = sample_auto_width_dataframe(dataframe)
    for column_index, column_name in enumerate(dataframe.columns, start=1):
        max_length = len(str(column_name))
        if column_name in sample.columns:
            lengths = sample[column_name].map(cell_display_length)
            if not lengths.empty:
                max_length = max(max_length, int(lengths.max()))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 40)


def sample_auto_width_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    if len(dataframe) <= AUTO_WIDTH_SAMPLE_LIMIT:
        return dataframe

    head_count = AUTO_WIDTH_SAMPLE_LIMIT // 2
    tail_count = AUTO_WIDTH_SAMPLE_LIMIT - head_count
    return pd.concat([dataframe.head(head_count), dataframe.tail(tail_count)], ignore_index=True)


def paths_refer_to_same_file(left_path: str | Path, right_path: str | Path) -> bool:
    left = Path(left_path)
    right = Path(right_path)
    try:
        left_marker = os.path.normcase(str(left.resolve()))
        right_marker = os.path.normcase(str(right.resolve()))
    except OSError:
        left_marker = os.path.normcase(str(left))
        right_marker = os.path.normcase(str(right))
    return left_marker == right_marker


__all__ = [
    "apply_auto_width",
    "apply_header_style",
    "clear_worksheet_values_preserve_structure",
    "export_dataframe",
    "export_dataframe_with_old_workbook",
    "export_to_excel",
    "paths_refer_to_same_file",
    "sample_auto_width_dataframe",
    "unmerge_ranges_overlapping_write_area",
    "write_dataframe_back_to_source",
    "write_dataframe_to_existing_excel_sheet",
    "write_rows_to_worksheet",
]
